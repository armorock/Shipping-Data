"""
Phase 1 — Job Code Markdown Generator
Reads the 'Job Code Data' tab from All Shipping Data BABY.xlsm and writes one
Obsidian markdown file per job code into FULL JOBCODE SHIPPING HISTORY/Job Codes/.
Covers the full AAA to max_code alphabetical sequence, with blank placeholders for
codes that have no Excel data.
"""

import os
import re
import openpyxl

# ── Paths ─────────────────────────────────────────────────────────────────────

EXCEL_PATH = (
    r"C:\Users\AlecSchooley\Desktop\Schooley Chaotic mind"
    r"\Projects\Shipping Data\All Shipping Data BABY.xlsm"
)
OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "Job Codes")
SHEET_NAME = "Job Code Data"
SEQ_START  = "AAA"

# ── Helpers ───────────────────────────────────────────────────────────────────

def code_to_int(code: str) -> int:
    n = 0
    for ch in code.upper():
        n = n * 26 + (ord(ch) - ord("A"))
    return n


def int_to_code(n: int) -> str:
    chars = []
    for _ in range(3):
        chars.append(chr(n % 26 + ord("A")))
        n //= 26
    return "".join(reversed(chars))


def sanitize_filename(name: str) -> str:
    """Strip characters illegal in Windows filenames."""
    return re.sub(r'[\\/*?:"<>|]', "", name).strip()


def extract_hyperlink_text(value) -> str:
    """Pull the display text out of an Excel HYPERLINK formula, or return raw."""
    if not value:
        return ""
    s = str(value)
    m = re.search(r'HYPERLINK\("[^"]*",\s*"([^"]+)"\)', s)
    if m:
        return m.group(1).strip()
    return s.strip()


def non_empty(row, limit: int) -> int:
    return sum(1 for v in row[:limit] if v is not None and str(v).strip())


def safe(value, cast=str) -> str:
    if value is None:
        return ""
    s = str(value).strip()
    if s.lower() in ("none", "nan"):
        return ""
    try:
        return str(cast(value)).strip()
    except Exception:
        return s


# ── Load Excel data ───────────────────────────────────────────────────────────

print(f"Reading {EXCEL_PATH!r} …")
wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, keep_vba=True)
ws = wb[SHEET_NAME]

# Col indices (0-based): Year=0, Job_Code=1, Project=2, County=3, ZipCode=4,
#   City=5, State=6, Customer=7, Amount=8, Plant=9, BOM_PDF=10, BOM_Link=11
HDR = {
    "year": 0, "code": 1, "project": 2, "county": 3, "zip": 4,
    "city": 5, "state": 6, "customer": 7, "amount": 8,
    "plant": 9, "bom_pdf": 10, "bom_link": 11,
}

raw_rows: list[dict] = []
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0:
        continue  # skip header
    code_raw = row[HDR["code"]]
    if not code_raw or not isinstance(code_raw, str):
        continue
    code = code_raw.strip().upper()
    if not (len(code) == 3 and code.isalpha()):
        continue
    raw_rows.append({"_row": i, "_filled": non_empty(row, 12), "code": code,
                     "year": row[HDR["year"]], "project": row[HDR["project"]],
                     "county": row[HDR["county"]], "zip": row[HDR["zip"]],
                     "city": row[HDR["city"]], "state": row[HDR["state"]],
                     "customer": row[HDR["customer"]], "amount": row[HDR["amount"]],
                     "plant": row[HDR["plant"]], "bom_pdf": row[HDR["bom_pdf"]],
                     "bom_link": row[HDR["bom_link"]]})

wb.close()
print(f"  {len(raw_rows)} valid rows read, {len(set(r['code'] for r in raw_rows))} unique codes.")

# ── Deduplicate: best row per code ────────────────────────────────────────────

lookup: dict[str, dict] = {}
for r in raw_rows:
    code = r["code"]
    existing = lookup.get(code)
    if existing is None:
        lookup[code] = r
    else:
        # Prefer more filled fields; break ties with higher year
        def year_val(x):
            try:
                return int(x["year"]) if x["year"] else 0
            except Exception:
                return 0
        if (r["_filled"], year_val(r)) > (existing["_filled"], year_val(existing)):
            lookup[code] = r

# ── Determine sequence end ────────────────────────────────────────────────────

max_code = max(lookup.keys(), key=code_to_int)
all_codes = [int_to_code(i) for i in range(code_to_int(SEQ_START), code_to_int(max_code) + 1)]
print(f"  Sequence: {SEQ_START} to {max_code}  ({len(all_codes)} total codes)")

# ── Write markdown files ──────────────────────────────────────────────────────

os.makedirs(OUTPUT_DIR, exist_ok=True)

created = 0
updated = 0
placeholder = 0

STATE_ABBREV = {
    "Alabama": "AL", "Alaska": "AK", "Arizona": "AZ", "Arkansas": "AR",
    "California": "CA", "Colorado": "CO", "Connecticut": "CT", "Delaware": "DE",
    "Florida": "FL", "Georgia": "GA", "Hawaii": "HI", "Idaho": "ID",
    "Illinois": "IL", "Indiana": "IN", "Iowa": "IA", "Kansas": "KS",
    "Kentucky": "KY", "Louisiana": "LA", "Maine": "ME", "Maryland": "MD",
    "Massachusetts": "MA", "Michigan": "MI", "Minnesota": "MN", "Mississippi": "MS",
    "Missouri": "MO", "Montana": "MT", "Nebraska": "NE", "Nevada": "NV",
    "New Hampshire": "NH", "New Jersey": "NJ", "New Mexico": "NM", "New York": "NY",
    "North Carolina": "NC", "North Dakota": "ND", "Ohio": "OH", "Oklahoma": "OK",
    "Oregon": "OR", "Pennsylvania": "PA", "Rhode Island": "RI", "South Carolina": "SC",
    "South Dakota": "SD", "Tennessee": "TN", "Texas": "TX", "Utah": "UT",
    "Vermont": "VT", "Virginia": "VA", "Washington": "WA", "West Virginia": "WV",
    "Wisconsin": "WI", "Wyoming": "WY", "District of Columbia": "DC",
}

TABLE_HEADER = (
    "## Shipped Items\n\n"
    "| Qty | Structure Name | Part Number | Part Type | Part Subtype"
    " | ID | Height | Opening | Description |\n"
    "| --- | -------------- | ----------- | --------- | ------------"
    " | -- | ------ | ------- | ----------- |\n"
)

for code in all_codes:
    data = lookup.get(code)

    if data:
        job_name    = safe(data["project"])
        customer    = safe(data["customer"])
        state       = safe(data["state"])
        county      = safe(data["county"])
        city        = safe(data["city"])
        zip_code    = safe(data["zip"])
        year_val    = safe(data["year"], int) if data["year"] else ""
        plant       = safe(data["plant"])
        bom_pdf     = extract_hyperlink_text(data["bom_pdf"])
    else:
        job_name = customer = state = county = city = ""
        zip_code = year_val = plant = bom_pdf = ""

    # YAML values: quote strings that need it
    def yaml_str(v: str) -> str:
        if not v:
            return ""
        if ":" in v or v != v.strip():
            return f'"{v}"'
        return v

    # Wiki links for state/county/city — quoted so YAML parses them as strings
    if state and state in STATE_ABBREV:
        abbrev = STATE_ABBREV[state]
    elif state:
        abbrev = state[:2].upper()
    else:
        abbrev = ""

    state_fm  = f'"[[{state}]]"'                                      if state  else ""
    county_fm = f'"[[{county}]]"'                                      if county else ""
    city_fm   = f'"[[{sanitize_filename(city + " " + abbrev)}]]"'     if city   else ""

    title = f"{code} - {job_name}" if job_name else code

    lines = [
        "---",
        f"job_code: {code}",
        f"customer: {yaml_str(customer)}",
        f"state: {state_fm}",
        f"county: {county_fm}",
        f"city: {city_fm}",
        f"zip: {zip_code}",
        f"year: {year_val}",
        f"plant: {yaml_str(plant)}",
        f"erp: ",
        f"bom_pdf: {yaml_str(bom_pdf)}",
        f"sharepoint_link: ",
        "---",
        "",
        f"# {title}",
        "",
        TABLE_HEADER,
    ]

    content = "\n".join(lines)

    safe_name = sanitize_filename(title)
    filename  = f"{safe_name}.md"
    filepath  = os.path.join(OUTPUT_DIR, filename)

    existed = os.path.exists(filepath)
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(content)

    if not data:
        placeholder += 1
    elif existed:
        updated += 1
    else:
        created += 1

print(f"\nDone.")
print(f"  Created  : {created}")
print(f"  Updated  : {updated}")
print(f"  Placeholders (no data): {placeholder}")
print(f"  Total files written   : {created + updated + placeholder}")
print(f"\nOutput to {OUTPUT_DIR}")
